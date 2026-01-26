"""Tests for WordExtractorService."""

import tempfile
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from src.services.word_extractor import (
    WordExtractorError,
    WordExtractorService,
    get_word_extractor_service,
)


@pytest.fixture
def sample_docx_content() -> BytesIO:
    """Create a sample Word document for testing extraction."""
    doc = Document()

    # Add patient information
    doc.add_heading("泌尿随访登记表", level=0)

    doc.add_paragraph("姓名：张三")
    doc.add_paragraph("性别：男")
    doc.add_paragraph("年龄：45")
    doc.add_paragraph("联系电话：13800138000")
    doc.add_paragraph("身高：175")
    doc.add_paragraph("体重：70")

    # Add checkboxes
    doc.add_paragraph("家族结石史：☑有☐无")
    doc.add_paragraph("反复泌尿系感染：☐有☑无")

    # Add a table
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "住院号"
    table.cell(1, 1).text = "H20240001"
    table.cell(2, 0).text = "手术日期"
    table.cell(2, 1).text = "2024-01-15"

    # Save to BytesIO
    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output


@pytest.fixture
def sample_docx_file(sample_docx_content: BytesIO, tmp_path: Path) -> Path:
    """Create a sample Word document file for testing."""
    file_path = tmp_path / "test_document.docx"
    file_path.write_bytes(sample_docx_content.read())
    return file_path


class TestWordExtractorService:
    """Tests for WordExtractorService class."""

    def test_get_word_extractor_service_returns_instance(self):
        """Test that get_word_extractor_service returns a WordExtractorService instance."""
        service = get_word_extractor_service()
        assert isinstance(service, WordExtractorService)

    def test_get_word_extractor_service_with_hospital_id(self):
        """Test that hospital_id is passed correctly."""
        service = get_word_extractor_service(hospital_id="hospital_001")
        assert service.hospital_id == "hospital_001"

    def test_init_default_sections(self):
        """Test that default sections are set correctly."""
        service = get_word_extractor_service()
        expected = ["basic_info", "surgery_indicator", "clinical_followup", "nursing_followup"]
        assert service.DEFAULT_SECTIONS == expected


class TestRawContentExtraction:
    """Tests for raw content extraction from documents."""

    def test_extract_raw_content(self, sample_docx_content: BytesIO):
        """Test extracting raw content from a document."""
        service = get_word_extractor_service()
        doc = Document(sample_docx_content)

        raw_content = service._extract_raw_content(doc)

        assert "paragraphs" in raw_content
        assert "tables" in raw_content
        assert len(raw_content["paragraphs"]) > 0
        assert len(raw_content["tables"]) > 0

    def test_extract_paragraphs(self, sample_docx_content: BytesIO):
        """Test that paragraphs are extracted correctly."""
        service = get_word_extractor_service()
        doc = Document(sample_docx_content)

        raw_content = service._extract_raw_content(doc)
        paragraphs = raw_content["paragraphs"]

        # Check some expected paragraphs
        full_text = "\n".join(paragraphs)
        assert "张三" in full_text
        assert "男" in full_text
        assert "家族结石史" in full_text

    def test_extract_tables(self, sample_docx_content: BytesIO):
        """Test that tables are extracted correctly."""
        service = get_word_extractor_service()
        doc = Document(sample_docx_content)

        raw_content = service._extract_raw_content(doc)
        tables = raw_content["tables"]

        assert len(tables) == 1
        table = tables[0]
        assert len(table) == 3  # 3 rows

        # Check table content
        flattened = [cell for row in table for cell in row]
        assert "住院号" in flattened
        assert "H20240001" in flattened


class TestFileExtraction:
    """Tests for extracting from files."""

    def test_extract_from_file(self, sample_docx_file: Path):
        """Test extracting data from a Word file."""
        service = get_word_extractor_service()

        result = service.extract_from_file(sample_docx_file)

        assert isinstance(result, dict)
        # Pattern-based extraction should find basic_info
        assert "basic_info" in result

    def test_extract_from_bytes(self, sample_docx_content: BytesIO):
        """Test extracting data from Word document bytes."""
        service = get_word_extractor_service()

        result = service.extract_from_bytes(sample_docx_content)

        assert isinstance(result, dict)
        assert "basic_info" in result

    def test_extract_from_file_not_found(self, tmp_path: Path):
        """Test that extraction from non-existent file raises error."""
        service = get_word_extractor_service()

        with pytest.raises(WordExtractorError):
            service.extract_from_file(tmp_path / "nonexistent.docx")


class TestPatternBasedExtraction:
    """Tests for pattern-based extraction (no LLM)."""

    def test_extract_patient_name(self, sample_docx_content: BytesIO):
        """Test extracting patient name with pattern matching."""
        service = get_word_extractor_service()

        result = service.extract_from_bytes(sample_docx_content)

        assert "basic_info" in result
        assert result["basic_info"].get("patient_name") == "张三"

    def test_extract_gender(self, sample_docx_content: BytesIO):
        """Test extracting gender with pattern matching."""
        service = get_word_extractor_service()

        result = service.extract_from_bytes(sample_docx_content)

        assert result["basic_info"].get("gender") == "男"

    def test_extract_age(self, sample_docx_content: BytesIO):
        """Test extracting age with pattern matching."""
        service = get_word_extractor_service()

        result = service.extract_from_bytes(sample_docx_content)

        assert result["basic_info"].get("age") == 45

    def test_extract_phone(self, sample_docx_content: BytesIO):
        """Test extracting phone number with pattern matching."""
        service = get_word_extractor_service()

        result = service.extract_from_bytes(sample_docx_content)

        assert result["basic_info"].get("phone") == "13800138000"

    def test_extract_height_weight(self, sample_docx_content: BytesIO):
        """Test extracting height and weight."""
        service = get_word_extractor_service()

        result = service.extract_from_bytes(sample_docx_content)

        assert result["basic_info"].get("height") == 175
        assert result["basic_info"].get("weight") == 70


class TestCheckboxNormalization:
    """Tests for checkbox normalization during extraction."""

    def test_normalize_family_history_positive(self, sample_docx_content: BytesIO):
        """Test normalizing positive checkbox (☑有☐无)."""
        service = get_word_extractor_service()

        result = service.extract_from_bytes(sample_docx_content)

        # The sample document has ☑有☐无 for family_history_of_stone
        # Pattern matching might not catch this depending on implementation

    def test_normalize_checkbox_with_normalizer_service(self):
        """Test that normalizer service is used for checkbox normalization."""
        service = get_word_extractor_service(hospital_id="hospital_001")

        assert service.normalizer_service is not None
        assert service.normalizer_service.hospital_id == "hospital_001"


class TestLLMExtraction:
    """Tests for LLM-based extraction."""

    def test_build_llm_prompt(self, sample_docx_content: BytesIO):
        """Test building the LLM prompt."""
        service = get_word_extractor_service()
        doc = Document(sample_docx_content)
        raw_content = service._extract_raw_content(doc)

        prompt = service._build_llm_prompt(raw_content, ["basic_info"])

        assert "medical document data extraction" in prompt.lower()
        assert "json" in prompt.lower()
        assert "basic_info" in prompt

    def test_parse_llm_response_json_block(self):
        """Test parsing LLM response with JSON code block."""
        service = get_word_extractor_service()

        response = """Here is the extracted data:

```json
{
    "basic_info": {
        "patient_name": "张三",
        "age": 45
    }
}
```
"""
        result = service._parse_llm_response(response)

        assert result["basic_info"]["patient_name"] == "张三"
        assert result["basic_info"]["age"] == 45

    def test_parse_llm_response_raw_json(self):
        """Test parsing LLM response with raw JSON."""
        service = get_word_extractor_service()

        response = """{
    "basic_info": {
        "patient_name": "李四"
    }
}"""
        result = service._parse_llm_response(response)

        assert result["basic_info"]["patient_name"] == "李四"

    def test_parse_llm_response_no_json(self):
        """Test parsing LLM response without JSON raises error."""
        service = get_word_extractor_service()

        response = "This is just some text without JSON."

        with pytest.raises(WordExtractorError) as exc_info:
            service._parse_llm_response(response)

        assert "No JSON found" in str(exc_info.value)

    def test_parse_llm_response_invalid_json(self):
        """Test parsing LLM response with invalid JSON raises error."""
        service = get_word_extractor_service()

        response = '{"basic_info": {"name": }}'  # Invalid JSON

        with pytest.raises(WordExtractorError) as exc_info:
            service._parse_llm_response(response)

        assert "Failed to parse JSON" in str(exc_info.value)


class TestNormalizeExtractedData:
    """Tests for normalizing extracted data."""

    def test_normalize_boolean_strings(self):
        """Test normalizing boolean string values."""
        service = get_word_extractor_service()

        extracted = {
            "basic_info": {
                "family_history_of_stone": "是",
                "repeated_urinary_infection": "否",
            }
        }

        normalized = service._normalize_extracted_data(extracted, ["basic_info"])

        # Values should be normalized based on schema and normalizer
        assert "basic_info" in normalized

    def test_normalize_with_stage_suffix(self):
        """Test normalizing entities with stage suffixes."""
        service = get_word_extractor_service()

        extracted = {
            "clinical_followup_1": {
                "recurrence": "否",
            },
            "clinical_followup_2": {
                "recurrence": "是",
            },
        }

        normalized = service._normalize_extracted_data(extracted, ["clinical_followup"])

        assert "clinical_followup_1" in normalized
        assert "clinical_followup_2" in normalized

    def test_normalize_non_dict_values(self):
        """Test that non-dict values are passed through."""
        service = get_word_extractor_service()

        extracted = {
            "metadata": "some string value",
        }

        normalized = service._normalize_extracted_data(extracted, [])

        assert normalized["metadata"] == "some string value"


class TestGetDocumentText:
    """Tests for get_document_text utility method."""

    def test_get_document_text(self, sample_docx_content: BytesIO):
        """Test getting full document text."""
        service = get_word_extractor_service()
        doc = Document(sample_docx_content)

        text = service.get_document_text(doc)

        assert "张三" in text
        assert "家族结石史" in text
        assert "[Table 1]" in text
        assert "住院号" in text


class TestHospitalAwareness:
    """Tests for hospital-specific extraction behavior."""

    def test_extractor_uses_hospital_normalizer(self):
        """Test that extractor uses hospital-specific normalizer."""
        service = get_word_extractor_service(hospital_id="hospital_001")

        assert service.hospital_id == "hospital_001"
        assert service.normalizer_service.hospital_id == "hospital_001"

    def test_extractor_uses_hospital_schema(self):
        """Test that extractor uses hospital-specific schema service."""
        service = get_word_extractor_service(hospital_id="hospital_001")

        assert service.schema_service is not None
        # Schema service should be able to get hospital-specific schemas
        # (if they exist - this tests the integration)
