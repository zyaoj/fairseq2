"""Tests for WordGeneratorService."""

import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from docx import Document

from src.services.word_generator import (
    WordGeneratorError,
    WordGeneratorService,
    get_word_generator_service,
)


@pytest.fixture
def template_dir(tmp_path: Path) -> Path:
    """Create a temporary template directory with a test template."""
    # Create base directory
    base_dir = tmp_path / "base"
    base_dir.mkdir()

    # Create a simple test template with placeholders
    doc = Document()
    doc.add_paragraph("Patient Name: {{basic_info.patient_name}}")
    doc.add_paragraph("Gender: {{basic_info.gender}}")
    doc.add_paragraph("Age: {{basic_info.age}}")
    doc.add_paragraph("Family History: {{basic_info.family_history_of_stone}}")
    doc.add_paragraph("Surgery Date: {{surgery_indicator.surgery_date}}")
    doc.add_paragraph("Surgery Method: {{surgery_indicator.surgery_method}}")
    doc.add_paragraph("Followup Date: {{clinical_followup_1.followup_date}}")

    # Add a table with placeholders
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Height"
    table.cell(0, 1).text = "{{basic_info.height}}"
    table.cell(1, 0).text = "Weight"
    table.cell(1, 1).text = "{{basic_info.weight}}"

    doc.save(base_dir / "followup_registration.docx")

    return tmp_path


class MockColumn:
    """Mock column object with a name attribute."""

    def __init__(self, name: str):
        self.name = name


class MockTable:
    """Mock table object with columns."""

    def __init__(self, column_names: list[str]):
        self.columns = [MockColumn(name) for name in column_names]


@pytest.fixture
def mock_patient() -> MagicMock:
    """Create a mock patient with related entities."""
    patient = MagicMock()
    patient.id = "test-patient-123"

    # Mock basic_info
    patient.basic_info = MagicMock()
    patient.basic_info.__table__ = MockTable([
        "id", "patient_id", "patient_name", "gender", "age",
        "height", "weight", "family_history_of_stone",
        "created_at", "updated_at",
    ])
    patient.basic_info.patient_name = "张三"
    patient.basic_info.gender = "男"
    patient.basic_info.age = 45
    patient.basic_info.height = 175.0
    patient.basic_info.weight = 70.0
    patient.basic_info.family_history_of_stone = True

    # Mock surgery_indicator
    patient.surgery_indicator = MagicMock()
    patient.surgery_indicator.__table__ = MockTable([
        "id", "patient_id", "surgery_date", "surgery_method",
        "created_at", "updated_at",
    ])
    patient.surgery_indicator.surgery_date = date(2024, 1, 15)
    patient.surgery_indicator.surgery_method = "经皮肾镜碎石术"

    # Mock clinical_followups
    clinical_followup = MagicMock()
    clinical_followup.stage = 1
    clinical_followup.__table__ = MockTable([
        "id", "patient_id", "stage", "followup_date", "recurrence",
        "created_at", "updated_at",
    ])
    clinical_followup.followup_date = date(2024, 1, 22)
    clinical_followup.recurrence = False
    patient.clinical_followups = [clinical_followup]

    # Mock nursing_followups
    patient.nursing_followups = []

    return patient


class TestWordGeneratorService:
    """Tests for WordGeneratorService class."""

    def test_get_word_generator_service_returns_instance(self):
        """Test that get_word_generator_service returns a WordGeneratorService instance."""
        service = get_word_generator_service()
        assert isinstance(service, WordGeneratorService)

    def test_get_word_generator_service_with_hospital_id(self):
        """Test that hospital_id is passed correctly."""
        service = get_word_generator_service(hospital_id="hospital_001")
        assert service.hospital_id == "hospital_001"

    def test_init_with_custom_template_service(self, template_dir: Path):
        """Test initialization with custom template service."""
        from src.services.template_service import TemplateService

        custom_service = TemplateService(template_dir=template_dir)
        generator = WordGeneratorService(template_service=custom_service)
        assert generator.template_service == custom_service


class TestDocumentGeneration:
    """Tests for document generation functionality."""

    def test_generate_followup_registration(self, template_dir: Path, mock_patient: MagicMock):
        """Test generating a followup registration document."""
        from src.services.template_service import TemplateService

        template_service = TemplateService(template_dir=template_dir)
        generator = WordGeneratorService(template_service=template_service)

        result = generator.generate_followup_registration(mock_patient)

        assert isinstance(result, BytesIO)
        assert result.tell() == 0  # Should be at the beginning

        # Verify the document can be opened
        doc = Document(result)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        assert "张三" in full_text
        assert "男" in full_text

    def test_generate_with_sections_filter(self, template_dir: Path, mock_patient: MagicMock):
        """Test generating document with specific sections."""
        from src.services.template_service import TemplateService

        template_service = TemplateService(template_dir=template_dir)
        generator = WordGeneratorService(template_service=template_service)

        result = generator.generate_followup_registration(
            mock_patient,
            include_sections=["basic_info"],
        )

        assert isinstance(result, BytesIO)

    def test_generate_with_missing_template_raises_error(self, tmp_path: Path, mock_patient: MagicMock):
        """Test that missing template raises WordGeneratorError."""
        from src.services.template_service import TemplateService

        # Create empty template directory
        (tmp_path / "base").mkdir()
        template_service = TemplateService(template_dir=tmp_path)
        generator = WordGeneratorService(template_service=template_service)

        with pytest.raises(WordGeneratorError):
            generator.generate_followup_registration(mock_patient)


class TestContextBuilding:
    """Tests for context building functionality."""

    def test_build_context_with_all_sections(self, template_dir: Path, mock_patient: MagicMock):
        """Test building context includes all sections."""
        from src.services.template_service import TemplateService

        template_service = TemplateService(template_dir=template_dir)
        generator = WordGeneratorService(template_service=template_service)

        context = generator._build_context(mock_patient)

        assert "basic_info" in context
        assert "surgery_indicator" in context
        assert "clinical_followup_1" in context

    def test_build_context_with_filtered_sections(self, template_dir: Path, mock_patient: MagicMock):
        """Test building context with specific sections only."""
        from src.services.template_service import TemplateService

        template_service = TemplateService(template_dir=template_dir)
        generator = WordGeneratorService(template_service=template_service)

        context = generator._build_context(
            mock_patient,
            include_sections=["basic_info"],
        )

        assert "basic_info" in context
        assert "surgery_indicator" not in context

    def test_build_context_with_missing_entity(self, template_dir: Path, mock_patient: MagicMock):
        """Test building context when entity is None."""
        from src.services.template_service import TemplateService

        mock_patient.basic_info = None

        template_service = TemplateService(template_dir=template_dir)
        generator = WordGeneratorService(template_service=template_service)

        context = generator._build_context(mock_patient)

        assert "basic_info" not in context


class TestValueFormatting:
    """Tests for value formatting functionality."""

    def test_format_boolean_yes_no(self):
        """Test formatting boolean values with yes/no format."""
        generator = get_word_generator_service()

        assert generator._format_boolean(True, "some_field") == "☑是☐否"
        assert generator._format_boolean(False, "some_field") == "☐是☑否"

    def test_format_boolean_has_none(self):
        """Test formatting boolean values with has/none format."""
        generator = get_word_generator_service()

        # Fields that should use has/none format
        assert generator._format_boolean(True, "family_history_of_stone") == "☑有☐无"
        assert generator._format_boolean(False, "family_history_of_stone") == "☐有☑无"
        assert generator._format_boolean(True, "repeated_urinary_infection") == "☑有☐无"

    def test_format_value_none(self):
        """Test formatting None values."""
        generator = get_word_generator_service()

        assert generator._format_value(None, "any_field") == ""

    def test_format_value_date(self):
        """Test formatting date values."""
        generator = get_word_generator_service()

        result = generator._format_value(date(2024, 1, 15), "surgery_date")
        assert result == "2024-01-15"

    def test_format_value_list(self):
        """Test formatting list values."""
        generator = get_word_generator_service()

        result = generator._format_value(["左肾", "右肾"], "stone_location")
        assert result == "左肾, 右肾"

    def test_format_value_float_integer(self):
        """Test formatting float values that are integers."""
        generator = get_word_generator_service()

        assert generator._format_value(175.0, "height") == "175"

    def test_format_value_float_decimal(self):
        """Test formatting float values with decimals."""
        generator = get_word_generator_service()

        assert generator._format_value(70.5, "weight") == "70.50"


class TestPlaceholderReplacement:
    """Tests for placeholder replacement functionality."""

    def test_placeholder_pattern_matches(self):
        """Test that placeholder pattern matches correctly."""
        generator = get_word_generator_service()

        matches = list(generator.PLACEHOLDER_PATTERN.finditer(
            "{{basic_info.patient_name}} and {{surgery_indicator.surgery_date}}"
        ))

        assert len(matches) == 2
        assert matches[0].group(1) == "basic_info"
        assert matches[0].group(2) == "patient_name"
        assert matches[1].group(1) == "surgery_indicator"
        assert matches[1].group(2) == "surgery_date"

    def test_placeholder_with_numbers(self):
        """Test placeholder pattern with numbered entities."""
        generator = get_word_generator_service()

        matches = list(generator.PLACEHOLDER_PATTERN.finditer(
            "{{clinical_followup_1.followup_date}}"
        ))

        assert len(matches) == 1
        assert matches[0].group(1) == "clinical_followup_1"
        assert matches[0].group(2) == "followup_date"
