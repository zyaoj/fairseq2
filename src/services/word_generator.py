"""Hospital-aware Word document generation service.

This service generates Word documents from patient data using hospital-specific templates.
It leverages the TemplateService for template resolution and supports the placeholder
syntax defined in the Word templates (e.g., {{basic_info.patient_name}}).
"""

import logging
import re
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

from docx import Document
from docx.table import Table, _Cell

from src.models.basic_info import BasicInfo
from src.models.clinical_followup import ClinicalFollowup
from src.models.nursing_followup import NursingFollowup
from src.models.patient import Patient
from src.models.surgery_indicator import SurgeryIndicator
from src.services.template_service import TemplateService, get_template_service


class WordGeneratorError(Exception):
    """Raised when Word document generation fails."""

    pass


class WordGeneratorService:
    """Service for generating Word documents from patient data.

    This service uses hospital-aware template resolution to generate Word documents
    with patient data filled in. Templates use placeholder syntax like:
    - {{basic_info.patient_name}}
    - {{surgery_indicator.stone_location}}
    - {{clinical_followup_1.followup_date}}
    - {{nursing_followup_2.nursing_mode}}

    Boolean fields are rendered as checkboxes (☑是/☐否 or ☑有/☐无).
    """

    # Placeholder pattern: {{entity.field}} - supports numbered entities like clinical_followup_1
    PLACEHOLDER_PATTERN = re.compile(r"\{\{([a-z_0-9]+)\.([a-z_0-9]+)\}\}")

    def __init__(
        self,
        template_service: TemplateService | None = None,
        hospital_id: str | None = None,
    ):
        """Initialize the Word generator service.

        Args:
            template_service: Template service for resolving templates.
                              Defaults to global template service.
            hospital_id: Hospital ID for hospital-specific templates.
        """
        self.template_service = template_service or get_template_service()
        self.hospital_id = hospital_id

    def generate_followup_registration(
        self,
        patient: Patient,
        include_sections: list[str] | None = None,
    ) -> BytesIO:
        """Generate a followup registration Word document for a patient.

        Args:
            patient: Patient with related entities (basic_info, surgery_indicator,
                     clinical_followups, nursing_followups)
            include_sections: Optional list of sections to include. If None, includes all.
                             Options: "basic_info", "surgery_indicator",
                                      "clinical_followups", "nursing_followups"

        Returns:
            BytesIO containing the generated Word document

        Raises:
            WordGeneratorError: If generation fails
        """
        try:
            # Get the template path
            template_path = self.template_service.get_template_path(
                "followup_registration", self.hospital_id
            )

            # Load the template
            doc = Document(template_path)

            # Build the data context from patient entities
            context = self._build_context(patient, include_sections)

            # Replace placeholders in the document
            self._replace_placeholders(doc, context)

            # Save to BytesIO
            output = BytesIO()
            doc.save(output)
            output.seek(0)

            return output

        except Exception as e:
            raise WordGeneratorError(f"Failed to generate document: {e}") from e

    def _build_context(
        self,
        patient: Patient,
        include_sections: list[str] | None = None,
    ) -> dict[str, dict[str, Any]]:
        """Build the data context from patient entities.

        Args:
            patient: Patient with related entities
            include_sections: Optional list of sections to include

        Returns:
            Dictionary mapping entity names to their field values
        """
        context: dict[str, dict[str, Any]] = {}

        # Default to all sections
        if include_sections is None:
            include_sections = [
                "basic_info",
                "surgery_indicator",
                "clinical_followups",
                "nursing_followups",
            ]

        # Basic info
        if "basic_info" in include_sections and patient.basic_info:
            context["basic_info"] = self._entity_to_dict(patient.basic_info)

        # Surgery indicator
        if "surgery_indicator" in include_sections and patient.surgery_indicator:
            context["surgery_indicator"] = self._entity_to_dict(patient.surgery_indicator)

        # Clinical followups (up to 5 stages)
        if "clinical_followups" in include_sections:
            for followup in patient.clinical_followups:
                key = f"clinical_followup_{followup.stage}"
                context[key] = self._entity_to_dict(followup)

        # Nursing followups (up to 6 stages)
        if "nursing_followups" in include_sections:
            for followup in patient.nursing_followups:
                key = f"nursing_followup_{followup.stage}"
                context[key] = self._entity_to_dict(followup)

        return context

    def _entity_to_dict(self, entity: Any) -> dict[str, Any]:
        """Convert a SQLAlchemy entity to a dictionary.

        Args:
            entity: SQLAlchemy model instance

        Returns:
            Dictionary of field names to values
        """
        result = {}

        # Get all columns from the entity
        for column in entity.__table__.columns:
            if column.name in ("id", "patient_id", "created_at", "updated_at"):
                continue  # Skip internal fields

            value = getattr(entity, column.name, None)
            result[column.name] = value

        return result

    def _replace_placeholders(
        self,
        doc: Document,
        context: dict[str, dict[str, Any]],
    ) -> None:
        """Replace placeholders in the document with actual values.

        Args:
            doc: Word document to modify
            context: Data context mapping entity.field to values
        """
        # Replace in paragraphs
        for paragraph in doc.paragraphs:
            self._replace_in_paragraph(paragraph, context)

        # Replace in tables
        for table in doc.tables:
            self._replace_in_table(table, context)

    def _replace_in_paragraph(
        self,
        paragraph: Any,
        context: dict[str, dict[str, Any]],
    ) -> None:
        """Replace placeholders in a paragraph.

        Args:
            paragraph: Word paragraph
            context: Data context
        """
        full_text = paragraph.text
        if "{{" not in full_text:
            return

        # Find all placeholders
        for match in self.PLACEHOLDER_PATTERN.finditer(full_text):
            entity_name = match.group(1)
            field_name = match.group(2)
            placeholder = match.group(0)

            # Get the value from context
            value = self._get_value(context, entity_name, field_name)
            formatted_value = self._format_value(value, field_name)

            # Replace in the full text
            full_text = full_text.replace(placeholder, formatted_value)

        # Update the paragraph text
        # We need to clear existing runs and add new text
        if full_text != paragraph.text:
            # Preserve formatting from first run if exists
            if paragraph.runs:
                first_run = paragraph.runs[0]
                # Clear all runs
                for run in paragraph.runs:
                    run.text = ""
                # Set new text on first run
                first_run.text = full_text
            else:
                paragraph.text = full_text

    def _replace_in_table(
        self,
        table: Table,
        context: dict[str, dict[str, Any]],
    ) -> None:
        """Replace placeholders in a table.

        Args:
            table: Word table
            context: Data context
        """
        for row in table.rows:
            for cell in row.cells:
                self._replace_in_cell(cell, context)

    def _replace_in_cell(
        self,
        cell: _Cell,
        context: dict[str, dict[str, Any]],
    ) -> None:
        """Replace placeholders in a table cell.

        Args:
            cell: Table cell
            context: Data context
        """
        for paragraph in cell.paragraphs:
            self._replace_in_paragraph(paragraph, context)

        # Handle nested tables
        for nested_table in cell.tables:
            self._replace_in_table(nested_table, context)

    def _get_value(
        self,
        context: dict[str, dict[str, Any]],
        entity_name: str,
        field_name: str,
    ) -> Any:
        """Get a value from the context.

        Args:
            context: Data context
            entity_name: Entity name (e.g., "basic_info", "clinical_followup_1")
            field_name: Field name (e.g., "patient_name", "followup_date")

        Returns:
            The field value or None if not found
        """
        entity_data = context.get(entity_name, {})
        return entity_data.get(field_name)

    def _format_value(self, value: Any, field_name: str) -> str:
        """Format a value for display in the Word document.

        Args:
            value: The raw value
            field_name: The field name (used to determine formatting)

        Returns:
            Formatted string value
        """
        if value is None:
            return ""

        # Format based on type
        if isinstance(value, bool):
            return self._format_boolean(value, field_name)
        elif isinstance(value, date):
            return value.strftime("%Y-%m-%d")
        elif isinstance(value, (list, tuple)):
            return ", ".join(str(v) for v in value)
        elif isinstance(value, float):
            # Format floats with 2 decimal places if they have decimals
            if value == int(value):
                return str(int(value))
            return f"{value:.2f}"
        else:
            return str(value)

    def _format_boolean(self, value: bool, field_name: str) -> str:
        """Format a boolean value as a checkbox.

        Args:
            value: Boolean value
            field_name: Field name (used to determine yes/no vs has/none format)

        Returns:
            Checkbox string (e.g., "☑是☐否" or "☐是☑否")
        """
        # Fields that use "有/无" (has/none) format
        has_none_fields = {
            "family_history_of_stone",
            "repeated_urinary_infection",
            "nursing_acidosis",
        }

        if field_name in has_none_fields:
            if value:
                return "☑有☐无"
            else:
                return "☐有☑无"
        else:
            # Default: yes/no format
            if value:
                return "☑是☐否"
            else:
                return "☐是☑否"


def get_word_generator_service(hospital_id: str | None = None) -> WordGeneratorService:
    """Get a Word generator service instance.

    Args:
        hospital_id: Optional hospital ID for hospital-specific templates

    Returns:
        WordGeneratorService instance
    """
    return WordGeneratorService(hospital_id=hospital_id)
