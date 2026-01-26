"""Hospital-aware Word document extraction service.

This service extracts structured data from Word documents using:
1. python-docx for text/table extraction
2. NormalizerService for hospital-specific checkbox normalization
3. Claude LLM for structured data extraction based on entity schemas
"""

import json
import logging
import re
from io import BytesIO
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

from docx import Document
from docx.table import Table

from src.services.normalizer import NormalizerService, get_normalizer_service
from src.services.schema_service import SchemaService, get_schema_service


class WordExtractorError(Exception):
    """Raised when Word document extraction fails."""

    pass


class WordExtractorService:
    """Service for extracting structured data from Word documents.

    This service parses Word documents and uses LLM to extract structured data
    following the entity schemas. It supports hospital-specific extraction rules
    via the NormalizerService.

    Extraction flow:
    1. Parse Word document text and tables
    2. Build LLM extraction prompt from hospital-merged schema
    3. Send to Claude LLM for structured data extraction
    4. Apply hospital-specific checkbox normalization
    5. Validate and return structured data
    """

    # Default sections to extract
    DEFAULT_SECTIONS = [
        "basic_info",
        "surgery_indicator",
        "clinical_followup",
        "nursing_followup",
    ]

    def __init__(
        self,
        schema_service: SchemaService | None = None,
        normalizer_service: NormalizerService | None = None,
        hospital_id: str | None = None,
        anthropic_client: Any = None,
    ):
        """Initialize the Word extractor service.

        Args:
            schema_service: Schema service for entity schemas.
            normalizer_service: Normalizer service for checkbox normalization.
            hospital_id: Hospital ID for hospital-specific rules.
            anthropic_client: Anthropic client for LLM extraction.
                              If None, extraction uses regex patterns only.
        """
        self.hospital_id = hospital_id
        self.schema_service = schema_service or get_schema_service()
        self.normalizer_service = normalizer_service or get_normalizer_service(
            hospital_id=hospital_id
        )
        self.anthropic_client = anthropic_client

    def extract_from_file(
        self,
        file_path: str | Path,
        sections: list[str] | None = None,
    ) -> dict[str, Any]:
        """Extract structured data from a Word document file.

        Args:
            file_path: Path to the Word document
            sections: Optional list of sections to extract

        Returns:
            Dictionary of extracted data by entity type

        Raises:
            WordExtractorError: If extraction fails
        """
        try:
            doc = Document(file_path)
            return self._extract_from_document(doc, sections)
        except Exception as e:
            raise WordExtractorError(f"Failed to extract from file: {e}") from e

    def extract_from_bytes(
        self,
        content: bytes | BytesIO,
        sections: list[str] | None = None,
    ) -> dict[str, Any]:
        """Extract structured data from Word document bytes.

        Args:
            content: Word document content as bytes or BytesIO
            sections: Optional list of sections to extract

        Returns:
            Dictionary of extracted data by entity type

        Raises:
            WordExtractorError: If extraction fails
        """
        try:
            if isinstance(content, bytes):
                content = BytesIO(content)
            doc = Document(content)
            return self._extract_from_document(doc, sections)
        except Exception as e:
            raise WordExtractorError(f"Failed to extract from bytes: {e}") from e

    def _extract_from_document(
        self,
        doc: Document,
        sections: list[str] | None = None,
    ) -> dict[str, Any]:
        """Extract structured data from a parsed Word document.

        Args:
            doc: Parsed Word document
            sections: Optional list of sections to extract

        Returns:
            Dictionary of extracted data by entity type
        """
        if sections is None:
            sections = self.DEFAULT_SECTIONS

        # Extract raw text and table content
        raw_content = self._extract_raw_content(doc)

        # If we have an LLM client, use it for extraction
        if self.anthropic_client:
            return self._extract_with_llm(raw_content, sections)

        # Otherwise, use pattern-based extraction
        return self._extract_with_patterns(raw_content, sections)

    def _extract_raw_content(self, doc: Document) -> dict[str, Any]:
        """Extract raw text and table content from the document.

        Args:
            doc: Parsed Word document

        Returns:
            Dictionary with 'paragraphs' and 'tables' keys
        """
        content = {
            "paragraphs": [],
            "tables": [],
        }

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content["paragraphs"].append(text)

        # Extract tables
        for table in doc.tables:
            table_data = self._extract_table(table)
            if table_data:
                content["tables"].append(table_data)

        return content

    def _extract_table(self, table: Table) -> list[list[str]]:
        """Extract table content as a 2D list.

        Args:
            table: Word table

        Returns:
            2D list of cell contents
        """
        rows = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                cells.append(cell_text)
            if any(cells):  # Only add rows with content
                rows.append(cells)
        return rows

    def _extract_with_llm(
        self,
        raw_content: dict[str, Any],
        sections: list[str],
    ) -> dict[str, Any]:
        """Extract structured data using Claude LLM.

        Args:
            raw_content: Raw document content
            sections: Sections to extract

        Returns:
            Dictionary of extracted data by entity type
        """
        # Build the prompt with schema information
        prompt = self._build_llm_prompt(raw_content, sections)

        # Call Claude API
        try:
            response = self.anthropic_client.messages.create(
                model="claude-sonnet-4-20250514",
                max_tokens=4096,
                messages=[
                    {"role": "user", "content": prompt}
                ],
            )

            # Parse the response
            response_text = response.content[0].text

            # Extract JSON from the response
            extracted = self._parse_llm_response(response_text)

            # Apply normalization to checkbox fields
            normalized = self._normalize_extracted_data(extracted, sections)

            return normalized

        except Exception as e:
            raise WordExtractorError(f"LLM extraction failed: {e}") from e

    def _build_llm_prompt(
        self,
        raw_content: dict[str, Any],
        sections: list[str],
    ) -> str:
        """Build the LLM extraction prompt.

        Args:
            raw_content: Raw document content
            sections: Sections to extract

        Returns:
            Prompt string for the LLM
        """
        # Get schema information for each section
        schema_info = []
        for section in sections:
            try:
                schema = self.schema_service.get_schema(section, self.hospital_id)
                fields = [
                    {
                        "name": f.name,
                        "type": f.field_type,
                        "description": f.i18n_key or f.name,
                    }
                    for f in schema.fields
                ]
                schema_info.append({
                    "entity": section,
                    "fields": fields,
                })
            except Exception:
                # Schema not found, skip
                continue

        # Build document content string
        doc_text = "\n".join(raw_content["paragraphs"])
        for i, table in enumerate(raw_content["tables"]):
            doc_text += f"\n\n[Table {i + 1}]\n"
            for row in table:
                doc_text += " | ".join(row) + "\n"

        prompt = f"""You are a medical document data extraction assistant. Extract structured data from the following Chinese urology follow-up registration document.

## Target Schema

Extract data for these entities and their fields:

{json.dumps(schema_info, ensure_ascii=False, indent=2)}

## Document Content

{doc_text}

## Instructions

1. Extract values for each field in the schema
2. For boolean fields, return true/false based on checkbox markers (☑是/☐否, ☑有/☐无, etc.)
3. For date fields, return in YYYY-MM-DD format
4. For list fields (like stone_location), return as JSON arrays
5. Return null for fields with no data

## Output Format

Return ONLY a JSON object with this structure:
{{
  "basic_info": {{...}},
  "surgery_indicator": {{...}},
  "clinical_followup_1": {{...}},
  "clinical_followup_2": {{...}},
  ...
  "nursing_followup_1": {{...}},
  ...
}}

Only include entities that have data in the document.
"""
        return prompt

    def _parse_llm_response(self, response_text: str) -> dict[str, Any]:
        """Parse the LLM response to extract JSON data.

        Args:
            response_text: LLM response text

        Returns:
            Extracted data dictionary
        """
        # Try to find JSON in the response
        # Look for JSON block markers first
        json_match = re.search(r"```json\s*([\s\S]*?)\s*```", response_text)
        if json_match:
            json_str = json_match.group(1)
        else:
            # Try to find raw JSON object
            json_match = re.search(r"\{[\s\S]*\}", response_text)
            if json_match:
                json_str = json_match.group(0)
            else:
                raise WordExtractorError("No JSON found in LLM response")

        try:
            return json.loads(json_str)
        except json.JSONDecodeError as e:
            raise WordExtractorError(f"Failed to parse JSON from LLM response: {e}") from e

    def _normalize_extracted_data(
        self,
        extracted: dict[str, Any],
        sections: list[str],
    ) -> dict[str, Any]:
        """Apply normalization to extracted checkbox fields.

        Args:
            extracted: Extracted data from LLM
            sections: Sections that were extracted

        Returns:
            Normalized data dictionary
        """
        normalized = {}

        for entity_name, entity_data in extracted.items():
            if not isinstance(entity_data, dict):
                normalized[entity_name] = entity_data
                continue

            # Determine the base entity type (strip stage numbers)
            base_entity = re.sub(r"_\d+$", "", entity_name)

            # Get schema for this entity
            try:
                schema = self.schema_service.get_schema(base_entity, self.hospital_id)
            except Exception:
                # No schema, just copy as-is
                normalized[entity_name] = entity_data
                continue

            # Normalize each field
            normalized_entity = {}
            for field_name, value in entity_data.items():
                # Find field in schema
                field = next(
                    (f for f in schema.fields if f.name == field_name),
                    None
                )

                if field and field.field_type == "boolean" and isinstance(value, str):
                    # Normalize boolean string to actual boolean
                    result = self.normalizer_service.normalize_yes_no(value)
                    normalized_entity[field_name] = result.normalized_value
                else:
                    normalized_entity[field_name] = value

            normalized[entity_name] = normalized_entity

        return normalized

    def _extract_with_patterns(
        self,
        raw_content: dict[str, Any],
        sections: list[str],
    ) -> dict[str, Any]:
        """Extract structured data using regex patterns (no LLM).

        This is a fallback method when no LLM client is available.
        It uses simple pattern matching to extract common fields.

        Args:
            raw_content: Raw document content
            sections: Sections to extract

        Returns:
            Dictionary of extracted data by entity type
        """
        extracted: dict[str, Any] = {}
        full_text = "\n".join(raw_content["paragraphs"])

        # Add table content to full text
        for table in raw_content["tables"]:
            for row in table:
                full_text += "\n" + " ".join(row)

        # Pattern-based extraction for common fields
        patterns = {
            "patient_name": r"姓名[：:]\s*(\S+)",
            "gender": r"性别[：:]\s*([男女])",
            "age": r"年龄[：:]\s*(\d+)",
            "phone": r"联系电话[：:]\s*(\d+)",
            "height": r"身高[：:]\s*(\d+(?:\.\d+)?)",
            "weight": r"体重[：:]\s*(\d+(?:\.\d+)?)",
        }

        basic_info = {}
        for field_name, pattern in patterns.items():
            match = re.search(pattern, full_text)
            if match:
                value = match.group(1)
                # Convert numeric fields
                if field_name in ("age", "height", "weight"):
                    try:
                        value = float(value) if "." in value else int(value)
                    except ValueError:
                        pass
                basic_info[field_name] = value

        # Extract boolean fields using normalizer
        bool_patterns = {
            "family_history_of_stone": r"家族结石史[：:]?\s*([☑☐√○✓✔][是否有无][☑☐√○✓✔][是否有无])",
            "repeated_urinary_infection": r"反复泌尿系感染[：:]?\s*([☑☐√○✓✔][是否有无][☑☐√○✓✔][是否有无])",
        }

        for field_name, pattern in bool_patterns.items():
            match = re.search(pattern, full_text)
            if match:
                checkbox_text = match.group(1)
                result = self.normalizer_service.normalize_yes_no(checkbox_text)
                if result.normalized_value is not None:
                    basic_info[field_name] = result.normalized_value

        if basic_info:
            extracted["basic_info"] = basic_info

        return extracted

    def get_document_text(self, doc: Document) -> str:
        """Get full text content from a Word document.

        Useful for debugging or manual inspection.

        Args:
            doc: Parsed Word document

        Returns:
            Full text content as string
        """
        raw_content = self._extract_raw_content(doc)
        text_parts = raw_content["paragraphs"]

        for i, table in enumerate(raw_content["tables"]):
            text_parts.append(f"\n[Table {i + 1}]")
            for row in table:
                text_parts.append(" | ".join(row))

        return "\n".join(text_parts)


def get_word_extractor_service(
    hospital_id: str | None = None,
    anthropic_client: Any = None,
) -> WordExtractorService:
    """Get a Word extractor service instance.

    Args:
        hospital_id: Optional hospital ID for hospital-specific rules
        anthropic_client: Optional Anthropic client for LLM extraction

    Returns:
        WordExtractorService instance
    """
    return WordExtractorService(
        hospital_id=hospital_id,
        anthropic_client=anthropic_client,
    )
