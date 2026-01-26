#!/usr/bin/env python3
"""Script to generate the base followup registration Word template."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def add_section_heading(doc: Document, text: str) -> None:
    """Add a section heading."""
    heading = doc.add_heading(text, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_subsection_heading(doc: Document, text: str) -> None:
    """Add a subsection heading."""
    doc.add_heading(text, level=2)


def add_info_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    """Add a simple two-column info table with label and placeholder."""
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for label, placeholder in rows:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = placeholder
        # Set cell widths
        row.cells[0].width = Inches(2)
        row.cells[1].width = Inches(4)


def add_checkbox_field(doc: Document, label: str, placeholder: str) -> None:
    """Add a checkbox-style field (used for yes/no fields)."""
    para = doc.add_paragraph()
    para.add_run(f"{label}: ").bold = True
    para.add_run(placeholder)


def add_multi_checkbox_field(
    doc: Document, label: str, options: list[tuple[str, str]]
) -> None:
    """Add a multi-checkbox field with multiple options."""
    para = doc.add_paragraph()
    para.add_run(f"{label}: ").bold = True
    for option_label, placeholder in options:
        para.add_run(f"{placeholder}{option_label}  ")


def create_followup_registration_template(output_path: Path) -> None:
    """Create the followup registration Word template."""
    doc = Document()

    # Title
    title = doc.add_heading("泌尿系感染随访登记表", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # =========================================================================
    # Section 1: Basic Info (基本信息)
    # =========================================================================
    add_section_heading(doc, "一、基本信息")

    basic_info_rows = [
        ("患者姓名", "{{basic_info.patient_name}}"),
        ("性别", "{{basic_info.gender}}"),
        ("年龄", "{{basic_info.age}}"),
        ("身高 (cm)", "{{basic_info.height}}"),
        ("体重 (kg)", "{{basic_info.weight}}"),
        ("职业", "{{basic_info.occupation}}"),
        ("联系电话", "{{basic_info.phone}}"),
    ]
    add_info_table(doc, basic_info_rows)
    doc.add_paragraph()

    add_subsection_heading(doc, "1.1 病史")

    # Checkbox fields for medical history
    add_checkbox_field(doc, "家族结石史", "{{basic_info.family_history_of_stone}}")
    add_checkbox_field(doc, "反复泌尿系感染", "{{basic_info.repeated_urinary_infection}}")
    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("首次泌尿系感染时间: ").bold = True
    para.add_run("{{basic_info.first_urinary_infection_time}}")

    para = doc.add_paragraph()
    para.add_run("既往泌尿系感染病原体: ").bold = True
    para.add_run("{{basic_info.previous_urinary_infection_pathogen}}")

    para = doc.add_paragraph()
    para.add_run("既往感染用药: ").bold = True
    para.add_run("{{basic_info.previous_infection_medication}}")

    para = doc.add_paragraph()
    para.add_run("既往病史: ").bold = True
    para.add_run("{{basic_info.medical_history}}")

    doc.add_paragraph()

    add_subsection_heading(doc, "1.2 生活方式")

    para = doc.add_paragraph()
    para.add_run("饮食偏好: ").bold = True
    para.add_run("{{basic_info.diet_preference}}")

    para = doc.add_paragraph()
    para.add_run("每日饮水量: ").bold = True
    para.add_run("{{basic_info.daily_water_intake}}")

    para = doc.add_paragraph()
    para.add_run("生活习惯: ").bold = True
    para.add_run("{{basic_info.lifestyle}}")

    doc.add_paragraph()

    add_subsection_heading(doc, "1.3 手术信息")

    para = doc.add_paragraph()
    para.add_run("手术日期: ").bold = True
    para.add_run("{{basic_info.surgery_date}}")

    para = doc.add_paragraph()
    para.add_run("手术方式: ").bold = True
    para.add_run("{{basic_info.surgery_type}}")

    para = doc.add_paragraph()
    para.add_run("结石成分: ").bold = True
    para.add_run("{{basic_info.stone_composition}}")

    doc.add_paragraph()

    # =========================================================================
    # Section 2: Surgery Indicators (手术指标)
    # =========================================================================
    add_section_heading(doc, "二、手术指标")

    para = doc.add_paragraph()
    para.add_run("临床诊断: ").bold = True
    para.add_run("{{surgery_indicator.clinical_diagnosis}}")

    doc.add_paragraph()

    # Stone location with multi-checkbox
    add_multi_checkbox_field(
        doc,
        "结石位置",
        [
            ("左肾", "{{surgery_indicator.stone_location.left_kidney}}"),
            ("右肾", "{{surgery_indicator.stone_location.right_kidney}}"),
            ("左输尿管", "{{surgery_indicator.stone_location.left_ureter}}"),
            ("右输尿管", "{{surgery_indicator.stone_location.right_ureter}}"),
            ("膀胱", "{{surgery_indicator.stone_location.bladder}}"),
        ],
    )

    para = doc.add_paragraph()
    para.add_run("结石大小: ").bold = True
    para.add_run("{{surgery_indicator.stone_size}}")

    para = doc.add_paragraph()
    para.add_run("肾积水程度: ").bold = True
    para.add_run("{{surgery_indicator.hydronephrosis_degree}}")

    doc.add_paragraph()

    # Pre-operative lab values
    add_subsection_heading(doc, "2.1 术前检验结果")
    preop_rows = [
        ("ALT", "{{surgery_indicator.alt_value_before}}"),
        ("AST", "{{surgery_indicator.ast_value_before}}"),
        ("GGT", "{{surgery_indicator.ggt_value_before}}"),
        ("血肌酐 (SCr)", "{{surgery_indicator.scr_value_before}}"),
        ("白细胞 (WBC)", "{{surgery_indicator.wbc_value_before}}"),
        ("血红蛋白 (Hb)", "{{surgery_indicator.hb_value_before}}"),
        ("尿pH", "{{surgery_indicator.urine_pH_value_before}}"),
        ("尿亚硝酸盐 (NIT)", "{{surgery_indicator.urine_NIT_value_before}}"),
        ("尿白细胞", "{{surgery_indicator.urine_WBC_value_before}}"),
        ("尿培养", "{{surgery_indicator.urine_culture_result_before}}"),
        ("尿NGS", "{{surgery_indicator.urine_NGS_result_before}}"),
    ]
    add_info_table(doc, preop_rows)
    doc.add_paragraph()

    # Post-operative lab values
    add_subsection_heading(doc, "2.2 术后检验结果")
    postop_rows = [
        ("ALT", "{{surgery_indicator.alt_value_after}}"),
        ("AST", "{{surgery_indicator.ast_value_after}}"),
        ("GGT", "{{surgery_indicator.ggt_value_after}}"),
        ("血肌酐 (SCr)", "{{surgery_indicator.scr_value_after}}"),
        ("白细胞 (WBC)", "{{surgery_indicator.wbc_value_after}}"),
        ("血红蛋白 (Hb)", "{{surgery_indicator.hb_value_after}}"),
        ("尿pH", "{{surgery_indicator.urine_pH_value_after}}"),
        ("尿亚硝酸盐 (NIT)", "{{surgery_indicator.urine_NIT_value_after}}"),
        ("尿白细胞", "{{surgery_indicator.urine_WBC_value_after}}"),
        ("尿培养", "{{surgery_indicator.urine_culture_result_after}}"),
        ("尿NGS", "{{surgery_indicator.urine_NGS_result_after}}"),
        ("结石培养", "{{surgery_indicator.stone_culture_result_after}}"),
        ("结石NGS", "{{surgery_indicator.stone_NGS_result_after}}"),
        ("结石成分分析", "{{surgery_indicator.stone_composition_after}}"),
        ("结石清除状态", "{{surgery_indicator.stone_clearance_after}}"),
        ("影像学检查方法", "{{surgery_indicator.imaging_method_after}}"),
    ]
    add_info_table(doc, postop_rows)
    doc.add_paragraph()

    # =========================================================================
    # Section 3: Clinical Follow-up (临床随访 - 5 stages)
    # =========================================================================
    add_section_heading(doc, "三、临床随访")

    clinical_stages = [
        (1, "术后7天"),
        (2, "术后1个月"),
        (3, "术后3个月"),
        (4, "术后6个月"),
        (5, "术后12个月"),
    ]

    for stage_num, stage_name in clinical_stages:
        add_subsection_heading(doc, f"3.{stage_num} {stage_name}")

        # Followup date
        para = doc.add_paragraph()
        para.add_run("随访日期: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.date}}}}")

        # Recurrence and stone status
        para = doc.add_paragraph()
        para.add_run("复发情况: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.recurrence}}}}")

        para = doc.add_paragraph()
        para.add_run("结石大小: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.stone_size}}}}")

        para = doc.add_paragraph()
        para.add_run("影像学结果: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.imaging}}}}")

        # Lab values table
        stage_lab_rows = [
            ("ALT", f"{{{{clinical_followup.stage_{stage_num}.alt}}}}"),
            ("AST", f"{{{{clinical_followup.stage_{stage_num}.ast}}}}"),
            ("GGT", f"{{{{clinical_followup.stage_{stage_num}.ggt}}}}"),
            ("血肌酐 (SCr)", f"{{{{clinical_followup.stage_{stage_num}.scr}}}}"),
            ("白细胞 (WBC)", f"{{{{clinical_followup.stage_{stage_num}.wbc}}}}"),
            ("血红蛋白 (Hb)", f"{{{{clinical_followup.stage_{stage_num}.hb}}}}"),
            ("尿pH", f"{{{{clinical_followup.stage_{stage_num}.urine_ph}}}}"),
            ("尿亚硝酸盐", f"{{{{clinical_followup.stage_{stage_num}.urine_nit}}}}"),
            ("尿白细胞", f"{{{{clinical_followup.stage_{stage_num}.urine_wbc}}}}"),
            ("尿培养", f"{{{{clinical_followup.stage_{stage_num}.urine_culture}}}}"),
        ]
        add_info_table(doc, stage_lab_rows)
        doc.add_paragraph()

        # Medication tracking
        para = doc.add_paragraph()
        para.add_run("用药情况: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.medication}}}}")

        para = doc.add_paragraph()
        para.add_run("用药剂量: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.medication_dose}}}}")

        para = doc.add_paragraph()
        para.add_run("用药天数: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.medication_days}}}}")

        para = doc.add_paragraph()
        para.add_run("依从性: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.compliance}}}}")

        para = doc.add_paragraph()
        para.add_run("不良反应: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.adverse}}}}")

        para = doc.add_paragraph()
        para.add_run("随访计划: ").bold = True
        para.add_run(f"{{{{clinical_followup.stage_{stage_num}.plan}}}}")

        doc.add_paragraph()

    # =========================================================================
    # Section 4: Nursing Follow-up (护理随访 - 6 stages)
    # =========================================================================
    add_section_heading(doc, "四、护理随访")

    nursing_stages = [
        (1, "术后7天"),
        (2, "术后1个月"),
        (3, "术后3个月"),
        (4, "术后6个月"),
        (5, "术后9个月"),
        (6, "术后12个月"),
    ]

    for stage_num, stage_name in nursing_stages:
        add_subsection_heading(doc, f"4.{stage_num} {stage_name}")

        # Followup mode and date
        para = doc.add_paragraph()
        para.add_run("随访方式: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.mode}}}}")

        para = doc.add_paragraph()
        para.add_run("随访日期: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.date}}}}")

        para = doc.add_paragraph()
        para.add_run("随访地点: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.location}}}}")

        doc.add_paragraph()

        # Medication section
        para = doc.add_paragraph()
        para.add_run("出院带药: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.discharge_drug}}}}")

        para = doc.add_paragraph()
        para.add_run("抗生素: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.antibiotic}}}}")

        para = doc.add_paragraph()
        para.add_run("抗生素剂量: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.antibiotic_dose}}}}")

        para = doc.add_paragraph()
        para.add_run("其他用药: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.other_drug}}}}")

        para = doc.add_paragraph()
        para.add_run("其他用药剂量: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.other_drug_dose}}}}")

        add_checkbox_field(
            doc,
            "定时服药",
            f"{{{{nursing_followup.stage_{stage_num}.timed_med}}}}",
        )

        doc.add_paragraph()

        # Clinical indicators
        para = doc.add_paragraph()
        para.add_run("尿pH: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.urine_ph}}}}")

        add_checkbox_field(
            doc,
            "酸中毒",
            f"{{{{nursing_followup.stage_{stage_num}.acidosis}}}}",
        )

        para = doc.add_paragraph()
        para.add_run("HCO3: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.hco3}}}}")

        para = doc.add_paragraph()
        para.add_run("不良反应: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.adverse_effects}}}}")

        doc.add_paragraph()

        # Hydration and urine
        para = doc.add_paragraph()
        para.add_run("饮水情况: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.water}}}}")

        add_checkbox_field(
            doc,
            "排尿量>2L",
            f"{{{{nursing_followup.stage_{stage_num}.urine_output}}}}",
        )

        para = doc.add_paragraph()
        para.add_run("尿色色卡等级 (0-8): ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.urine_color}}}}")

        para = doc.add_paragraph()
        para.add_run("残石排出: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.residual_stone}}}}")

        doc.add_paragraph()

        # Lifestyle
        para = doc.add_paragraph()
        para.add_run("饮食指导: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.diet}}}}")

        doc.add_paragraph()

        # Mental health
        para = doc.add_paragraph()
        para.add_run("PSQI评分: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.psqi}}}}")

        para = doc.add_paragraph()
        para.add_run("睡眠用药: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.psqi_drug}}}}")

        para = doc.add_paragraph()
        para.add_run("抑郁筛查: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.depression}}}}")

        para = doc.add_paragraph()
        para.add_run("焦虑抑郁评估: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.depression_anxiety}}}}")

        para = doc.add_paragraph()
        para.add_run("心理疏导: ").bold = True
        para.add_run(f"{{{{nursing_followup.stage_{stage_num}.support}}}}")

        doc.add_paragraph()

    # =========================================================================
    # Section 5: Signature
    # =========================================================================
    doc.add_paragraph()
    add_section_heading(doc, "五、签名")

    para = doc.add_paragraph()
    para.add_run("医生签名: ").bold = True
    para.add_run("__________________")
    para.add_run("    日期: ").bold = True
    para.add_run("__________________")

    doc.add_paragraph()

    para = doc.add_paragraph()
    para.add_run("护士签名: ").bold = True
    para.add_run("__________________")
    para.add_run("    日期: ").bold = True
    para.add_run("__________________")

    # Save the document
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Template created: {output_path}")


if __name__ == "__main__":
    # Create the base template
    base_template_path = (
        Path(__file__).parent.parent
        / "src"
        / "word_templates"
        / "base"
        / "followup_registration.docx"
    )
    create_followup_registration_template(base_template_path)
